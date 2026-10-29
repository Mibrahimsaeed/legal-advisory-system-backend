import io
import json

import pytest
from docx import Document as DocxDocument

from config.settings import get_settings


async def _chat(client, headers, question="How do I apply for bail?"):
    conv = (await client.post("/conversations", json={"title": "Bail question"}, headers=headers)).json()
    answer = (
        await client.post(f"/chat/{conv['id']}/message", json={"question": question}, headers=headers)
    ).json()
    return conv, answer


async def test_save_answer_as_pdf_and_download(client, auth_headers):
    conv, answer = await _chat(client, auth_headers)
    r = await client.post(
        "/documents/generated",
        json={"source": "answer", "conversation_id": conv["id"], "message_id": answer["id"], "format": "pdf"},
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["kind"] == "answer"
    assert doc["file_format"] == "pdf"
    assert doc["source_title"] == "Bail question"
    assert doc["filename"].endswith(".pdf")

    d = await client.get(f"/documents/generated/{doc['id']}/download", headers=auth_headers)
    assert d.status_code == 200
    assert d.headers["content-type"] == "application/pdf"
    assert d.content.startswith(b"%PDF")

    listed = (await client.get("/documents/generated", headers=auth_headers)).json()
    assert any(x["id"] == doc["id"] for x in listed)


async def test_save_conversation_as_word(client, auth_headers):
    conv, _ = await _chat(client, auth_headers, "What is khula?")
    r = await client.post(
        "/documents/generated",
        json={"source": "conversation", "conversation_id": conv["id"], "format": "docx"},
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text
    d = await client.get(f"/documents/generated/{r.json()['id']}/download", headers=auth_headers)
    text = "\n".join(p.text for p in DocxDocument(io.BytesIO(d.content)).paragraphs)
    assert "What is khula?" in text
    assert "not legal advice" in text


async def test_answer_requires_message_id_and_assistant_message(client, auth_headers):
    conv, _ = await _chat(client, auth_headers)
    r = await client.post(
        "/documents/generated", json={"source": "answer", "conversation_id": conv["id"], "format": "pdf"},
        headers=auth_headers,
    )
    assert r.status_code == 422
    r = await client.post(
        "/documents/generated",
        json={"source": "answer", "conversation_id": conv["id"], "message_id": "nope", "format": "pdf"},
        headers=auth_headers,
    )
    assert r.status_code == 404


async def test_other_user_cannot_download_or_delete(client, auth_headers):
    conv, answer = await _chat(client, auth_headers)
    doc = (await client.post(
        "/documents/generated",
        json={"source": "answer", "conversation_id": conv["id"], "message_id": answer["id"], "format": "pdf"},
        headers=auth_headers,
    )).json()

    await client.post("/auth/register", json={"email": "other.gen@example.com", "password": "otherpass123"})
    tok = (await client.post("/auth/login", json={"email": "other.gen@example.com", "password": "otherpass123"})).json()
    other = {"Authorization": f"Bearer {tok['access_token']}"}
    assert (await client.get(f"/documents/generated/{doc['id']}/download", headers=other)).status_code == 403
    assert (await client.delete(f"/documents/generated/{doc['id']}", headers=other)).status_code == 403
    assert (await client.get(f"/documents/generated/{doc['id']}/download")).status_code == 401

    assert (await client.delete(f"/documents/generated/{doc['id']}", headers=auth_headers)).status_code == 204
    assert (await client.get(f"/documents/generated/{doc['id']}/download", headers=auth_headers)).status_code == 404


@pytest.fixture
def case_law_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "case_law_data_dir", str(tmp_path))
    return tmp_path


async def test_case_law_not_available_without_data(client, auth_headers, case_law_dir):
    s = (await client.get("/documents/generated/case-law/status", headers=auth_headers)).json()
    assert s == {"available": False, "cases": 0}
    r = await client.post("/documents/generated/case-law", json={"query": "khula", "format": "pdf"}, headers=auth_headers)
    assert r.status_code == 404
    assert "not available" in r.json()["detail"].lower()


async def test_case_law_uses_only_cases_in_the_data(client, auth_headers, case_law_dir):
    # Obviously fake test records, used only inside this test.
    (case_law_dir / "cases.json").write_text(json.dumps([
        {"citation": "TEST 2000 XX 1", "title": "Test Party A v Test Party B", "court": "Test Court",
         "headnote": "Test headnote about khula and dissolution of marriage.", "judgment": "Test judgment text."},
        {"citation": "TEST 2000 XX 2", "case_title": "Test Party C v Test Party D",
         "summary": "Test summary about bail after arrest."},
        {"title": "Record without any text is ignored"},
    ]))
    s = (await client.get("/documents/generated/case-law/status", headers=auth_headers)).json()
    assert s == {"available": True, "cases": 2}

    r = await client.post("/documents/generated/case-law", json={"query": "khula", "format": "docx"}, headers=auth_headers)
    assert r.status_code == 201, r.text
    d = await client.get(f"/documents/generated/{r.json()['id']}/download", headers=auth_headers)
    text = "\n".join(p.text for p in DocxDocument(io.BytesIO(d.content)).paragraphs)
    assert "TEST 2000 XX 1" in text
    assert "TEST 2000 XX 2" not in text

    r = await client.post("/documents/generated/case-law", json={"query": "income tax appeal", "format": "pdf"}, headers=auth_headers)
    assert r.status_code == 404
    assert "not available" in r.json()["detail"].lower()

    r = await client.post("/documents/generated/case-law", json={"query": "TEST 2000 XX 2", "format": "pdf"}, headers=auth_headers)
    assert r.status_code == 201
